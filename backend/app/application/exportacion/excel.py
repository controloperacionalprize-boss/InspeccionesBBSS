"""Generación de archivos Excel (.xlsx) para las exportaciones de listados."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

CONTENT_TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _valor_celda(valor: Any) -> Any:
    if isinstance(valor, (date, datetime)):
        return valor.isoformat()
    return valor if valor is not None else ""


def libro_excel(hoja: str, encabezados: list[str], filas: list[list[Any]]) -> bytes:
    """Construye un .xlsx en memoria con una sola hoja y devuelve sus bytes."""
    libro = Workbook()
    ws = libro.active
    ws.title = hoja[:31]  # Excel limita el nombre de hoja a 31 caracteres.

    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True)

    for fila in filas:
        ws.append([_valor_celda(valor) for valor in fila])

    for indice, encabezado in enumerate(encabezados, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = max(12, min(40, len(str(encabezado)) + 4))

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def nombre_archivo(prefijo: str) -> str:
    return f"{prefijo}_{date.today().isoformat()}.xlsx"
